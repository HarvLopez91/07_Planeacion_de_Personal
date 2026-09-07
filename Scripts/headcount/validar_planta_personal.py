#!/usr/bin/env python3
"""Valida la hoja `Planta Personal` de PptovsReal.xlsx contra Consolidado 2025.xlsx.

SOLO LECTURA: abre los libros con `read_only=True` y nunca llama a `save()`.
No emite datos personales: todas las salidas son totales agregados.
El reporte se escribe en `Outputs/`, ignorado por Git.

Uso:
    python Scripts/headcount/validar_planta_personal.py [ANIO] [NN.Mes]
    python Scripts/headcount/validar_planta_personal.py 2026 08.Agosto
"""
import collections
import datetime
import io
import os
import sys
import unicodedata

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ORIGEN = os.path.join(RAIZ, "Data", "HeadCount", "2025", "Consolidado 2025.xlsx")
DESTINO = os.path.join(RAIZ, "Data", "HeadCount", "PptovsReal.xlsx")
SALIDA = os.path.join(RAIZ, "Outputs")

APRENDIZ = {"SENA", "CONTRATO APRENDIZAJE", "CONTRATO DE APRENDIZAJE"}
FIJO = {"FIJO", "CONTRATO FIJO"}
INDEFINIDO = {"INDEFINIDO", "CONTRATO INDEFINIDO"}
TEMPORAL = {"TEMPORAL"}


def norm(v):
    """Normaliza a mayusculas sin acentos para comparar nombres de empresa."""
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def leer_origen():
    """Agrega Consolidado2025 por (anio, mes, grupo, empresa) aplicando la homologacion."""
    wb = openpyxl.load_workbook(ORIGEN, read_only=True, data_only=True)
    ws = wb["Consolidado2025"]
    filas = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else "" for c in next(filas)]
    i_grupo, i_emp = 0, 1
    i_tipo = cab.index("TIPO_CONTR")
    i_mes = cab.index("MES")
    i_anio = i_mes + 1  # la cabecera de anio trae acento inconsistente entre versiones
    agg = collections.defaultdict(collections.Counter)
    for r in filas:
        if r[i_emp] is None:
            continue
        clave = (r[i_anio], str(r[i_mes]).strip(), norm(r[i_grupo]), norm(r[i_emp]))
        c = agg[clave]
        t = str(r[i_tipo]).strip().upper()
        c["Total"] += 1
        if t in APRENDIZ:
            c["Sena"] += 1
        elif t in FIJO:
            c["Fijos"] += 1
        elif t in INDEFINIDO:
            c["Indefinidos"] += 1
        elif t in TEMPORAL:
            c["Temporales"] += 1
        else:
            c["SinClasificar"] += 1
    wb.close()
    return agg


def leer_destino():
    """Lee las filas `Real` de la hoja Planta Personal."""
    wb = openpyxl.load_workbook(DESTINO, read_only=True, data_only=True)
    ws = wb["Planta Personal"]
    filas = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else "" for c in next(filas)]
    ix = {n: k for k, n in enumerate(cab)}
    i_anio = next((v for k, v in ix.items() if k.startswith("A") and k.endswith("o")), 1)
    out = []
    for r in filas:
        if r[ix["Ppto/Real"]] is None:
            continue
        if str(r[ix["Ppto/Real"]]).strip().upper() != "REAL":
            continue
        out.append({
            "anio": r[i_anio],
            "mes": str(r[ix["Mes"]]).strip(),
            "grupo": norm(r[ix["Grupo Empresa"]]),
            "empresa": norm(r[ix["Empresa"]]),
            "Fijos": r[ix["Fijos"]],
            "Indefinidos": r[ix["Indefinidos"]],
            "Sena": r[ix["Sena"]],
            "Temporales": r[ix["Temporales"]],
            "Total": r[ix["Total"]],
            "Total-Sena": r[ix["Total-Sena"]],
        })
    wb.close()
    return out


def validar(anio, mes):
    origen = leer_origen()
    destino = leer_destino()
    lineas = []
    fallos = []

    def chk(ok, titulo, detalle=""):
        lineas.append("- [%s] %s%s" % ("PASS" if ok else "FALLA", titulo,
                                       ("  -- " + detalle) if detalle else ""))
        if not ok:
            fallos.append(titulo)

    per = [d for d in destino if d["anio"] == anio and d["mes"] == mes]
    ori = {k: v for k, v in origen.items() if k[0] == anio and k[1] == mes}
    con_dot = [d for d in per if d["Total"] is not None]

    chk(bool(con_dot), "1. El periodo existe como Real con dotacion",
        "%d filas, %d con dotacion" % (len(per), len(con_dot)))

    emp_o = set(k[3] for k in ori)
    emp_d = set(d["empresa"] for d in con_dot)
    faltan = sorted(emp_o - emp_d)
    chk(not faltan, "2. Todas las empresas del origen estan en el destino",
        ("faltan: " + ", ".join(faltan)) if faltan else "%d empresas" % len(emp_o))

    extra = sorted(emp_d - emp_o)
    detalle_extra = ", ".join(
        "%s=%s" % (e, next(d["Total"] for d in con_dot if d["empresa"] == e)) for e in extra)
    en_cero = [e for e in extra
               if next(d["Total"] for d in con_dot if d["empresa"] == e) == 0]
    chk(extra == sorted(en_cero), "3. Empresas sin registros en el origen figuran en cero",
        detalle_extra or "ninguna")

    total_o = sum(c["Total"] for c in ori.values())
    ts_o = sum(c["Total"] - c["Sena"] for c in ori.values())
    total_d = sum(d["Total"] or 0 for d in con_dot)
    ts_d = sum(d["Total-Sena"] or 0 for d in con_dot)
    chk(total_o == total_d, "4a. Total del periodo cuadra",
        "origen %d vs destino %d" % (total_o, total_d))
    chk(ts_o == ts_d, "4b. Total-Sena del periodo cuadra",
        "origen %d vs destino %d" % (ts_o, ts_d))

    malas = [d["empresa"] for d in con_dot
             if (d["Fijos"] or 0) + (d["Indefinidos"] or 0) + (d["Sena"] or 0)
             + (d["Temporales"] or 0) != (d["Total"] or 0)
             or (d["Total"] or 0) - (d["Sena"] or 0) != (d["Total-Sena"] or 0)]
    chk(not malas, "5. Coherencia interna: suma de contratos y Total-Sena",
        ("filas incoherentes: " + ", ".join(malas)) if malas else "%d filas" % len(con_dot))

    dest_ix = {}
    for d in destino:
        if d["Total"] is not None:
            dest_ix[(d["anio"], d["mes"], d["empresa"])] = d["Total"]
    difs = []
    for (a, m, _g, e), c in origen.items():
        clave = (a, m, e)
        if clave in dest_ix and dest_ix[clave] != c["Total"]:
            difs.append("%s %s %s: origen %d vs destino %d"
                        % (a, m, e, c["Total"], dest_ix[clave]))
    chk(not difs, "6. Sin diferencias por empresa en ningun periodo",
        ("%d diferencias" % len(difs)) if difs else "todos los periodos comparables coinciden")

    llaves = collections.Counter(
        (d["anio"], d["mes"], d["grupo"], d["empresa"]) for d in destino)
    dup = [k for k, n in llaves.items() if n > 1]
    chk(not dup, "7. Sin duplicados por Anio/Mes/Ppto-Real/Grupo/Empresa",
        ("%d llaves duplicadas" % len(dup)) if dup else "%d llaves unicas" % len(llaves))

    return lineas, fallos, difs, con_dot


def main():
    anio = int(sys.argv[1]) if len(sys.argv) > 1 else 2026
    mes = sys.argv[2] if len(sys.argv) > 2 else "08.Agosto"
    lineas, fallos, difs, per = validar(anio, mes)
    estado = "PASS" if not fallos else "FALLA"

    rep = ["# Validacion de `Planta Personal` - %s %s" % (anio, mes), "",
           "Fecha: %s" % datetime.date.today().isoformat(),
           "Resultado: **%s**" % estado, "",
           "Solo lectura. Sin datos personales: unicamente totales agregados.", "",
           "## Validaciones", ""]
    rep += lineas
    rep += ["", "## Totales del periodo", "",
            "| Grupo | Empresa | Fijos | Indef | Sena | Temp | Total | Total-Sena |",
            "|---|---|---:|---:|---:|---:|---:|---:|"]
    for d in sorted(per, key=lambda x: (x["grupo"], x["empresa"])):
        rep.append("| %s | %s | %s | %s | %s | %s | %s | %s |"
                   % (d["grupo"], d["empresa"], d["Fijos"], d["Indefinidos"],
                      d["Sena"], d["Temporales"], d["Total"], d["Total-Sena"]))
    rep.append("| **TOTAL** | | %d | %d | %d | %d | **%d** | **%d** |"
               % (sum(d["Fijos"] or 0 for d in per),
                  sum(d["Indefinidos"] or 0 for d in per),
                  sum(d["Sena"] or 0 for d in per),
                  sum(d["Temporales"] or 0 for d in per),
                  sum(d["Total"] or 0 for d in per),
                  sum(d["Total-Sena"] or 0 for d in per)))
    if difs:
        rep += ["", "## Diferencias por empresa", ""] + ["- " + d for d in difs]

    os.makedirs(SALIDA, exist_ok=True)
    ruta = os.path.join(SALIDA, "validacion_planta_personal_%s_%s.md"
                        % (anio, mes.split(".")[0]))
    io.open(ruta, "w", encoding="utf-8", newline="").write("\r\n".join(rep))

    print("\n".join(lineas))
    print()
    print("RESULTADO: %s" % estado)
    print("Reporte: %s" % os.path.relpath(ruta, RAIZ))
    return 0 if estado == "PASS" else 1


if __name__ == "__main__":
    sys.exit(main())
