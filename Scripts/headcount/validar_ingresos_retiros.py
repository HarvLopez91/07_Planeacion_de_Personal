#!/usr/bin/env python3
"""Valida INGRESOS y RETIROS de PptovsReal contra el consolidador Kactus.

El proceso es estrictamente de solo lectura para los libros de entrada. Las
identificaciones se normalizan y comparan solo en memoria; nunca se incluyen en
la consola ni en el reporte agregado.

Codigos de salida:
    0: PASS (puede incluir advertencias no bloqueantes)
    1: inconsistencia de datos
    2: error de estructura o contrato
    3: error tecnico de lectura/configuracion
"""

from __future__ import annotations

import argparse
import base64
import collections
import datetime as dt
import json
import math
import re
import sys
import unicodedata
import zlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DESTINATION = ROOT / "Data" / "HeadCount" / "PptovsReal.xlsx"
DEFAULT_KACTUS = (
    ROOT
    / "Data"
    / "Contratos_Kactus"
    / "Fuente_Oficial"
    / "CONSOLIDADOR_CONTRATOS_V0.0.0.xlsx"
)
DEFAULT_EMPRESAS = (
    ROOT
    / "PBIP"
    / "Proyecto.SemanticModel"
    / "definition"
    / "tables"
    / "Empresas.tmdl"
)
DEFAULT_GRUPOS = DEFAULT_EMPRESAS.with_name("Grupo Empresarial.tmdl")
DEFAULT_OUTPUT = ROOT / "Outputs"

MONTH_ABBR = {
    1: "ene",
    2: "feb",
    3: "mar",
    4: "abr",
    5: "may",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "sep",
    10: "oct",
    11: "nov",
    12: "dic",
}
MONTH_NAMES = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "setiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}

DESTINATION_SCHEMA = {
    "year": ("Año", "Ano"),
    "month": ("Mes",),
    "group": ("Grupo Empresarial", "Grupo empresarial", "Grupo Empresa"),
    "company": ("Empresa",),
    "identifier": ("Identificacion", "Identificación"),
}
DESTINATION_DATE_COLUMNS = {
    "INGRESOS": ("Fecha Inicio",),
    "RETIROS": ("Fecha Vencimiento",),
}
KACTUS_ID_COLUMNS = ("Identificacion", "Identificación")
KACTUS_DATE_COLUMNS = {
    "INGRESOS": ("Fecha Inicio", "Fecha Contrato"),
    "RETIROS": ("Fecha Vencimiento",),
}


class ContractError(Exception):
    """El archivo existe, pero no cumple el contrato estructural."""


class TechnicalError(Exception):
    """La entrada no puede abrirse o la configuracion no es utilizable."""


@dataclass(frozen=True, order=True)
class Period:
    year: int
    month: int

    @classmethod
    def parse(cls, year: int | str, month: int | str) -> "Period":
        parsed_year = int(year)
        parsed_month = parse_month(month)
        if parsed_year < 1900 or parsed_year > 9999:
            raise ValueError("El año debe tener cuatro digitos validos")
        return cls(parsed_year, parsed_month)

    @property
    def key(self) -> str:
        return f"{self.year:04d}-{self.month:02d}"


@dataclass(frozen=True)
class Catalog:
    company_to_group: dict[str, str]
    groups: frozenset[str]


@dataclass(frozen=True)
class Record:
    group: str
    company: str
    identifier: str
    date: dt.date | None


@dataclass
class SheetData:
    records: list[Record]
    periods: dict[Period, set[str]] = field(default_factory=dict)
    invalid_dates: int = 0
    other_period_dates: int = 0
    invalid_period_values: int = 0
    source_names: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Finding:
    level: str
    category: str
    check: str
    detail: str = ""


@dataclass
class ValidationResult:
    period: Period
    findings: list[Finding] = field(default_factory=list)
    summaries: dict[str, collections.Counter] = field(default_factory=dict)
    counts: dict[str, tuple[int, int]] = field(default_factory=dict)

    def add(self, level: str, category: str, check: str, detail: str = "") -> None:
        self.findings.append(Finding(level, category, check, detail))

    @property
    def has_errors(self) -> bool:
        return any(f.level == "ERROR" for f in self.findings)

    @property
    def state(self) -> str:
        return "ERROR" if self.has_errors else "PASS"


def normalize_text(value: object) -> str:
    if isinstance(value, float) and math.isnan(value):
        value = None
    text = "" if value is None else str(value).strip()
    normalized = unicodedata.normalize("NFKD", text)
    return "".join(char for char in normalized if not unicodedata.combining(char)).casefold()


def normalize_identifier(value: object) -> str:
    """Normaliza una identificacion solo para comparacion privada en memoria."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# Alias conservado para compatibilidad con el delta local auditado.
clave_id = normalize_identifier


def parse_month(value: object) -> int:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.month
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        month = int(value)
        if 1 <= month <= 12:
            return month
    text = normalize_text(value)
    match = re.match(r"^(\d{1,2})(?:\D|$)", text)
    if match and 1 <= int(match.group(1)) <= 12:
        return int(match.group(1))
    for name, month in MONTH_NAMES.items():
        if name in text:
            return month
    raise ValueError(f"Mes no reconocido: {value!r}")


def parse_year(value: object) -> int:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.year
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    match = re.search(r"\b(19|20)\d{2}\b", str(value or ""))
    if not match:
        raise ValueError(f"Año no reconocido: {value!r}")
    return int(match.group(0))


def parse_date(value: object) -> dt.date | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def clean_kactus_header(value: object) -> str:
    text = str(value or "").strip()
    if "[" in text:
        text = text.rsplit("[", 1)[-1].rstrip("]")
    return text


def column_index(headers: Sequence[object], aliases: Iterable[str], label: str) -> int:
    normalized = {normalize_text(value): index for index, value in enumerate(headers)}
    for alias in aliases:
        if normalize_text(alias) in normalized:
            return normalized[normalize_text(alias)]
    raise ContractError(
        f"Falta columna obligatoria {label!r}; aceptadas: {', '.join(aliases)}"
    )


def require_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise TechnicalError(f"No existe {label}: {path}")


def decode_static_table(path: Path) -> list[list[object]]:
    require_file(path, "archivo TMDL")
    text = path.read_text(encoding="utf-8", errors="replace")
    match = re.search(r'Binary\.FromText\("([^"]+)", BinaryEncoding\.Base64\)', text)
    if not match:
        raise ContractError(f"No se encontro el catalogo embebido en {path.name}")
    try:
        payload = zlib.decompress(base64.b64decode(match.group(1)), -15)
        return json.loads(payload.decode("utf-8"))
    except (ValueError, zlib.error, json.JSONDecodeError) as exc:
        raise ContractError(f"Catalogo embebido invalido en {path.name}") from exc


def dimension_empresas(empresas_path: Path = DEFAULT_EMPRESAS) -> set[str]:
    """API compatible: devuelve las empresas validas de la dimension."""
    return {
        str(row[0]).strip()
        for row in decode_static_table(Path(empresas_path))
        if row and str(row[0]).strip()
    }


def load_catalog(empresas_path: Path, grupos_path: Path) -> Catalog:
    companies = dimension_empresas(empresas_path)
    groups = {str(row[0]).strip() for row in decode_static_table(grupos_path) if row}
    text = empresas_path.read_text(encoding="utf-8", errors="replace")
    pairs = re.findall(r'if \[Empresas\] = "([^"]+)" then "([^"]+)"', text)
    mapping: dict[str, str] = {}
    conflicting = set()
    for company, group in pairs:
        if company in mapping and mapping[company] != group:
            conflicting.add(company)
        mapping[company] = group
    if not companies or not groups or not mapping:
        raise ContractError("El catalogo Empresa/Grupo esta vacio o no pudo interpretarse")
    if conflicting:
        raise ContractError(
            "El TMDL contiene asignaciones de grupo contradictorias para una empresa"
        )
    # Desktop conserva dos pasos M equivalentes para esta columna; repetir el
    # mismo par es valido, pero cada empresa debe quedar cubierta una sola vez
    # en el mapa resultante.
    if set(mapping) != companies:
        raise ContractError("El mapeo Empresa/Grupo no cubre exactamente el catalogo de empresas")
    if not set(mapping.values()).issubset(groups):
        raise ContractError("El mapeo Empresa/Grupo contiene grupos fuera del catalogo oficial")
    return Catalog(mapping, frozenset(groups))


def catalogo_grupos_empresas(
    empresas_path: Path = DEFAULT_EMPRESAS,
    grupos_path: Path = DEFAULT_GRUPOS,
) -> tuple[dict[str, str], set[str]]:
    """API compatible: devuelve Empresa -> Grupo y grupos oficiales."""
    catalog = load_catalog(Path(empresas_path), Path(grupos_path))
    return dict(catalog.company_to_group), set(catalog.groups)


def open_workbook(path: Path) -> openpyxl.Workbook:
    require_file(path, "libro de entrada")
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise TechnicalError(f"No fue posible abrir {path}: {exc}") from exc


def destination_rows(
    path: Path,
    sheet_name: str,
    period: Period | None,
) -> SheetData:
    workbook = open_workbook(path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ContractError(f"Falta hoja obligatoria {sheet_name!r} en {path.name}")
        iterator = workbook[sheet_name].iter_rows(values_only=True)
        try:
            headers = next(iterator)
        except StopIteration as exc:
            raise ContractError(f"La hoja {sheet_name!r} esta vacia") from exc
        indexes = {
            key: column_index(headers, aliases, f"{sheet_name}.{key}")
            for key, aliases in DESTINATION_SCHEMA.items()
        }
        indexes["date"] = column_index(
            headers, DESTINATION_DATE_COLUMNS[sheet_name], f"{sheet_name}.fecha"
        )
        result = SheetData(records=[])
        for row in iterator:
            if not any(value not in (None, "") for value in row):
                continue
            year_value = row[indexes["year"]]
            month_value = row[indexes["month"]]
            if year_value in (None, "") and month_value in (None, ""):
                continue
            try:
                row_period = Period(
                    parse_year(year_value), parse_month(month_value)
                )
            except (ValueError, TypeError):
                result.invalid_period_values += 1
                continue
            group = str(row[indexes["group"]] or "").strip()
            result.periods.setdefault(row_period, set()).add(group)
            if period is not None and row_period != period:
                continue
            parsed_date = parse_date(row[indexes["date"]])
            if parsed_date is None:
                result.invalid_dates += 1
            elif Period(parsed_date.year, parsed_date.month) != row_period:
                result.other_period_dates += 1
            result.records.append(
                Record(
                    group=group,
                    company=str(row[indexes["company"]] or "").strip(),
                    identifier=normalize_identifier(row[indexes["identifier"]]),
                    date=parsed_date,
                )
            )
        return result
    finally:
        workbook.close()


def discover_latest_period(path: Path) -> Period:
    incomes = destination_rows(path, "INGRESOS", None).periods
    departures = destination_rows(path, "RETIROS", None).periods
    common = sorted(set(incomes) & set(departures))
    if not common:
        raise ContractError("INGRESOS y RETIROS no tienen un periodo comun")
    return common[-1]


def kactus_sheet_pattern(kind: str, period: Period) -> re.Pattern[str]:
    month = MONTH_ABBR[period.month]
    if kind == "INGRESOS":
        pattern = rf"^\s*{month}\s*,\s*A\s*$"
    else:
        pattern = rf"^\s*{month}\s*(?:II|,\s*I)\s*$"
    return re.compile(pattern, re.IGNORECASE)


def kactus_rows(path: Path, kind: str, period: Period) -> SheetData:
    workbook = open_workbook(path)
    try:
        pattern = kactus_sheet_pattern(kind, period)
        names = [name for name in workbook.sheetnames if pattern.match(name)]
        if not names:
            raise ContractError(
                f"No existe hoja Kactus para {kind} {period.key}; patron {pattern.pattern}"
            )
        result = SheetData(records=[], source_names=names)
        for name in names:
            rows = workbook[name].iter_rows(values_only=True)
            first_rows = []
            for _ in range(3):
                try:
                    first_rows.append(next(rows))
                except StopIteration as exc:
                    raise ContractError(f"La hoja Kactus {name!r} no contiene encabezados") from exc
            headers = [clean_kactus_header(value) for value in first_rows[2]]
            id_index = column_index(headers, KACTUS_ID_COLUMNS, f"{name}.identificacion")
            date_index = column_index(headers, KACTUS_DATE_COLUMNS[kind], f"{name}.fecha")
            for row in rows:
                if not any(value not in (None, "") for value in row):
                    continue
                parsed_date = parse_date(row[date_index])
                if parsed_date is None:
                    result.invalid_dates += 1
                elif Period(parsed_date.year, parsed_date.month) != period:
                    result.other_period_dates += 1
                result.records.append(
                    Record("", "", normalize_identifier(row[id_index]), parsed_date)
                )
        return result
    finally:
        workbook.close()


def duplicate_count(identifiers: Iterable[str]) -> int:
    counts = collections.Counter(identifier for identifier in identifiers if identifier)
    return sum(count - 1 for count in counts.values() if count > 1)


def validate_kind(
    result: ValidationResult,
    kind: str,
    destination: SheetData,
    source: SheetData,
    catalog: Catalog,
) -> None:
    target = destination.records
    origin = source.records
    result.counts[kind] = (len(target), len(origin))
    result.summaries[kind] = collections.Counter(
        (record.group, record.company) for record in target
    )

    def check(ok: bool, name: str, detail: str) -> None:
        result.add("PASS" if ok else "ERROR", "DATOS", f"{kind}: {name}", detail)

    check(bool(target), "periodo presente", f"{len(target)} registros")
    check(
        len(target) == len(origin),
        "total concilia con Kactus",
        f"PptovsReal={len(target)}; Kactus={len(origin)}; hojas={','.join(source.source_names)}",
    )
    missing_target = sum(not record.identifier for record in target)
    missing_origin = sum(not record.identifier for record in origin)
    check(
        missing_target == 0 and missing_origin == 0,
        "identificacion presente",
        f"vacias en PptovsReal={missing_target}; vacias en Kactus={missing_origin}",
    )
    duplicates = duplicate_count(record.identifier for record in target)
    check(duplicates == 0, "sin duplicados", f"duplicados adicionales={duplicates}")
    check(
        destination.invalid_dates == 0 and destination.other_period_dates == 0,
        "fechas validas en PptovsReal",
        f"invalidas={destination.invalid_dates}; fuera del periodo={destination.other_period_dates}",
    )
    check(
        destination.invalid_period_values == 0,
        "Año y Mes validos en PptovsReal",
        f"filas con periodo invalido={destination.invalid_period_values}",
    )
    check(
        source.invalid_dates == 0 and source.other_period_dates == 0,
        "fechas validas en Kactus",
        f"invalidas={source.invalid_dates}; fuera del periodo={source.other_period_dates}",
    )
    target_ids = collections.Counter(record.identifier for record in target)
    origin_ids = collections.Counter(record.identifier for record in origin)
    check(
        "" not in target_ids and "" not in origin_ids and target_ids == origin_ids,
        "registros concilian por identificacion",
        "faltantes=%d; adicionales=%d"
        % (
            sum((origin_ids - target_ids).values()),
            sum((target_ids - origin_ids).values()),
        ),
    )

    unknown_companies = sorted(
        {record.company for record in target if record.company not in catalog.company_to_group}
    )
    check(
        not unknown_companies,
        "Empresa homologada al catalogo oficial",
        f"sin correspondencia={len(unknown_companies)}"
        + (f" ({', '.join(unknown_companies[:5])})" if unknown_companies else ""),
    )
    invalid_groups = sum(
        catalog.company_to_group.get(record.company) != record.group for record in target
    )
    not_derivable = sum(
        catalog.company_to_group.get(record.company) not in catalog.groups for record in target
    )
    if kind == "RETIROS":
        check(
            not_derivable == 0,
            "Grupo efectivo derivable por Empresa",
            f"sin grupo oficial={not_derivable}",
        )
        if invalid_groups:
            result.add(
                "WARN",
                "DATOS",
                "RETIROS: columna Grupo empresarial redundante",
                f"discrepancias={invalid_groups}; el grupo efectivo se deriva por Empresa",
            )
    else:
        check(
            invalid_groups == 0,
            "Grupo empresarial homologado",
            f"registros incompatibles={invalid_groups}",
        )

    prior = sorted(period for period in destination.periods if period < result.period)
    if prior:
        previous = prior[-1]
        current_groups = {record.group for record in target}
        new_groups = sorted(current_groups - destination.periods[previous])
        result.add(
            "INFO",
            "DIAGNOSTICO",
            f"{kind}: grupos nuevos frente a {previous.key}",
            f"cantidad={len(new_groups)}"
            + (f" ({', '.join(new_groups)})" if new_groups else ""),
        )


def validate_files(
    destination_path: Path,
    kactus_path: Path,
    catalog: Catalog,
    period: Period,
) -> ValidationResult:
    result = ValidationResult(period)
    for kind in ("INGRESOS", "RETIROS"):
        destination = destination_rows(destination_path, kind, period)
        source = kactus_rows(kactus_path, kind, period)
        validate_kind(result, kind, destination, source, catalog)
    return result


def render_console(result: ValidationResult) -> str:
    lines = []
    for finding in result.findings:
        detail = f" -- {finding.detail}" if finding.detail else ""
        lines.append(f"[{finding.level}][{finding.category}] {finding.check}{detail}")
    lines.append("")
    lines.append(f"RESULTADO: {result.state}")
    return "\n".join(lines)


def render_report(result: ValidationResult) -> str:
    lines = [
        f"# Validacion de INGRESOS y RETIROS — {result.period.key}",
        "",
        f"Fecha de ejecucion: {dt.date.today().isoformat()}",
        f"Resultado: **{result.state}**",
        "",
        "Solo lectura. El reporte contiene totales agregados y nunca identificaciones.",
        "",
        "## Validaciones",
        "",
    ]
    for finding in result.findings:
        detail = f" — {finding.detail}" if finding.detail else ""
        lines.append(f"- **{finding.level} / {finding.category}** — {finding.check}{detail}")
    for kind in ("INGRESOS", "RETIROS"):
        target, origin = result.counts.get(kind, (0, 0))
        lines.extend(
            [
                "",
                f"## {kind}",
                "",
                f"PptovsReal: **{target}** | Kactus: **{origin}**",
                "",
                "| Grupo empresarial | Empresa | Registros |",
                "|---|---|---:|",
            ]
        )
        for (group, company), count in sorted(result.summaries.get(kind, {}).items()):
            lines.append(f"| {group} | {company} | {count} |")
        lines.append(f"| **TOTAL** | | **{target}** |")
    return "\r\n".join(lines) + "\r\n"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Valida PptovsReal contra Kactus sin modificar los libros."
    )
    parser.add_argument("year", nargs="?", help="Año; si se omite, descubre el ultimo comun")
    parser.add_argument("month", nargs="?", help="Mes numerico o NN.Nombre")
    parser.add_argument("--pptovsreal", type=Path, default=DEFAULT_DESTINATION)
    parser.add_argument("--kactus", type=Path, default=DEFAULT_KACTUS)
    parser.add_argument("--empresas-tmdl", type=Path, default=DEFAULT_EMPRESAS)
    parser.add_argument("--grupos-tmdl", type=Path, default=DEFAULT_GRUPOS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--no-report", action="store_true", help="No escribe reporte agregado")
    return parser


def resolve_period(args: argparse.Namespace) -> Period:
    if args.year is None and args.month is None:
        return discover_latest_period(args.pptovsreal)
    if args.year is None or args.month is None:
        raise ContractError("Año y mes deben indicarse juntos")
    try:
        return Period.parse(args.year, args.month)
    except (TypeError, ValueError) as exc:
        raise ContractError(str(exc)) from exc


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        period = resolve_period(args)
        catalog = load_catalog(args.empresas_tmdl, args.grupos_tmdl)
        result = validate_files(args.pptovsreal, args.kactus, catalog, period)
        print(render_console(result))
        if not args.no_report:
            args.output_dir.mkdir(parents=True, exist_ok=True)
            report_path = args.output_dir / f"validacion_ingresos_retiros_{period.key}.md"
            report_path.write_text(render_report(result), encoding="utf-8", newline="")
            print(f"Reporte: {report_path}")
        return 1 if result.has_errors else 0
    except ContractError as exc:
        print(f"[ERROR][ESTRUCTURA] {exc}", file=sys.stderr)
        return 2
    except TechnicalError as exc:
        print(f"[ERROR][TECNICO] {exc}", file=sys.stderr)
        return 3
    except (OSError, PermissionError) as exc:
        print(f"[ERROR][TECNICO] {exc}", file=sys.stderr)
        return 3


if __name__ == "__main__":
    sys.exit(main())
