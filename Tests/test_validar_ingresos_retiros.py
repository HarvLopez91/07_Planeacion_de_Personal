from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Scripts" / "headcount"))

import validar_ingresos_retiros as validator  # noqa: E402


CATALOG = validator.Catalog(
    company_to_group={
        "Empresa Uno": "Grupo Uno",
        "Empresa Dos": "Grupo Dos",
    },
    groups=frozenset({"Grupo Uno", "Grupo Dos"}),
)


def record(identifier: str, company: str = "Empresa Uno", group: str = "Grupo Uno"):
    return {"id": identifier, "company": company, "group": group}


def create_destination(path: Path, periods: dict[validator.Period, list[dict]]) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    headers = [
        "Grupo Empresarial",
        "Empresa",
        "Identificacion",
        "Año",
        "Mes",
        "Fecha Inicio",
        "Fecha Vencimiento",
    ]
    for kind in ("INGRESOS", "RETIROS"):
        sheet = workbook.create_sheet(kind)
        sheet.append(headers)
        for period, rows in sorted(periods.items()):
            for item in rows:
                date = item.get("date", dt.date(period.year, period.month, 10))
                sheet.append(
                    [
                        item.get("group", "Grupo Uno"),
                        item.get("company", "Empresa Uno"),
                        item.get("id"),
                        period.year,
                        f"{period.month:02d}.Mes",
                        date,
                        date,
                    ]
                )
    workbook.save(path)


def create_kactus(path: Path, period: validator.Period, rows: list[dict]) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for kind, name in (
        ("INGRESOS", f"{validator.MONTH_ABBR[period.month]}, A"),
        ("RETIROS", f"{validator.MONTH_ABBR[period.month].title()} II"),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append([f"Volcado {kind}"])
        sheet.append([])
        sheet.append(
            [
                "Fact_Contrataciones[Identificación]",
                "Fact_Contrataciones[Fecha Inicio]",
                "Fact_Contrataciones[Fecha Vencimiento]",
            ]
        )
        for item in rows:
            date = item.get("date", dt.date(period.year, period.month, 10))
            sheet.append([item.get("id"), date, date])
    workbook.save(path)


def validate(tmp_path: Path, rows: list[dict], period: validator.Period | None = None):
    period = period or validator.Period(2026, 8)
    destination = tmp_path / "PptovsReal.xlsx"
    kactus = tmp_path / "Kactus.xlsx"
    create_destination(destination, {period: rows})
    create_kactus(kactus, period, rows)
    return validator.validate_files(destination, kactus, CATALOG, period)


def findings(result, level: str):
    return [finding for finding in result.findings if finding.level == level]


def test_ingreso_y_retiro_validos_dan_pass(tmp_path):
    result = validate(tmp_path, [record("TEST-ID-001"), record("TEST-ID-002")])

    assert result.state == "PASS"
    assert not findings(result, "ERROR")
    assert result.counts == {"INGRESOS": (2, 2), "RETIROS": (2, 2)}


def test_duplicado_es_error_sin_publicar_identificacion(tmp_path):
    secret = "TEST-DOCUMENTO-SECRETO"
    result = validate(tmp_path, [record(secret), record(secret)])

    assert result.state == "ERROR"
    assert any("sin duplicados" in item.check for item in findings(result, "ERROR"))
    assert secret not in validator.render_console(result)
    assert secret not in validator.render_report(result)


def test_identificacion_faltante_es_error(tmp_path):
    result = validate(tmp_path, [record("")])

    assert result.state == "ERROR"
    assert any("identificacion presente" in item.check for item in findings(result, "ERROR"))


def test_fecha_invalida_es_error(tmp_path):
    result = validate(tmp_path, [dict(record("TEST-ID-001"), date="fecha-invalida")])

    assert result.state == "ERROR"
    assert any("fechas validas" in item.check for item in findings(result, "ERROR"))


def test_empresa_desconocida_es_error(tmp_path):
    result = validate(
        tmp_path,
        [record("TEST-ID-001", company="Empresa Nueva", group="Grupo Uno")],
    )

    assert result.state == "ERROR"
    assert any("Empresa homologada" in item.check for item in findings(result, "ERROR"))


def test_grupo_desconocido_bloquea_ingresos_y_advierte_en_retiros(tmp_path):
    result = validate(tmp_path, [record("TEST-ID-001", group="Grupo Incorrecto")])

    assert result.state == "ERROR"
    assert any("INGRESOS: Grupo empresarial" in item.check for item in findings(result, "ERROR"))
    assert any("RETIROS: columna Grupo empresarial" in item.check for item in findings(result, "WARN"))


def test_retiros_con_grupo_redundante_distinto_es_warning_no_bloqueante():
    period = validator.Period(2026, 8)
    destination = validator.SheetData(
        records=[validator.Record("Grupo Incorrecto", "Empresa Uno", "TEST-ID-001", dt.date(2026, 8, 1))]
    )
    source = validator.SheetData(
        records=[validator.Record("", "", "TEST-ID-001", dt.date(2026, 8, 1))],
        source_names=["Ago II"],
    )
    result = validator.ValidationResult(period)

    validator.validate_kind(result, "RETIROS", destination, source, CATALOG)

    assert not findings(result, "ERROR")
    assert any(item.level == "WARN" for item in result.findings)


def test_columna_obligatoria_ausente_es_error_de_contrato(tmp_path):
    path = tmp_path / "incompleto.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "INGRESOS"
    workbook.active.append(["Empresa"])
    workbook.create_sheet("RETIROS").append(["Empresa"])
    workbook.save(path)

    with pytest.raises(validator.ContractError, match="Falta columna obligatoria"):
        validator.destination_rows(path, "INGRESOS", validator.Period(2026, 8))


def test_archivo_vacio_es_error_de_contrato(tmp_path):
    path = tmp_path / "vacio.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "INGRESOS"
    workbook.save(path)

    with pytest.raises(validator.ContractError):
        validator.destination_rows(path, "INGRESOS", validator.Period(2026, 8))


def test_multiples_periodos_seleccionan_solo_el_solicitado(tmp_path):
    path = tmp_path / "PptovsReal.xlsx"
    july = validator.Period(2026, 7)
    august = validator.Period(2026, 8)
    create_destination(
        path,
        {
            july: [record("TEST-JULIO")],
            august: [record("TEST-AGOSTO-1"), record("TEST-AGOSTO-2")],
        },
    )

    data = validator.destination_rows(path, "INGRESOS", august)

    assert len(data.records) == 2
    assert set(data.periods) == {july, august}
    assert validator.discover_latest_period(path) == august


def test_ingesta_futura_funciona_sin_cambiar_codigo(tmp_path):
    future = validator.Period(2027, 1)
    result = validate(tmp_path, [record("TEST-FUTURO")], future)

    assert result.state == "PASS"
    assert result.period == future


def test_periodo_de_fecha_distinto_es_error(tmp_path):
    result = validate(
        tmp_path,
        [dict(record("TEST-ID-001"), date=dt.date(2026, 7, 31))],
        validator.Period(2026, 8),
    )

    assert result.state == "ERROR"
    assert any("fuera del periodo=1" in item.detail for item in findings(result, "ERROR"))


def test_nan_se_trata_como_identificacion_vacia():
    assert validator.normalize_identifier(float("nan")) == ""


def test_codigos_de_salida_distinguen_contrato_y_tecnico(tmp_path, capsys):
    absent = tmp_path / "no-existe.xlsx"
    code = validator.main(
        [
            "2026",
            "08.Agosto",
            "--pptovsreal",
            str(absent),
            "--kactus",
            str(absent),
            "--no-report",
        ]
    )

    assert code == 3
    assert "[ERROR]" in capsys.readouterr().err


def test_codigo_uno_distingue_inconsistencia_de_datos(tmp_path, monkeypatch):
    period = validator.Period(2026, 8)
    destination = tmp_path / "PptovsReal.xlsx"
    kactus = tmp_path / "Kactus.xlsx"
    duplicate_rows = [record("TEST-DUP"), record("TEST-DUP")]
    create_destination(destination, {period: duplicate_rows})
    create_kactus(kactus, period, duplicate_rows)
    monkeypatch.setattr(validator, "load_catalog", lambda *_: CATALOG)

    code = validator.main(
        [
            "2026",
            "08.Agosto",
            "--pptovsreal",
            str(destination),
            "--kactus",
            str(kactus),
            "--no-report",
        ]
    )

    assert code == 1
